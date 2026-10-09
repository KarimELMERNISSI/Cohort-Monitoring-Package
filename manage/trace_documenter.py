
import os
import json
from datetime import datetime
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

class TraceDocumenter:
    def __init__(self, trace_source):
        """
        Initialize with either a dictionary trace or a path to a json trace file.
        """
        self.trace = {}
        if isinstance(trace_source, dict):
            self.trace = trace_source
        elif isinstance(trace_source, str) and os.path.exists(trace_source):
            try:
                with open(trace_source, 'r') as f:
                    self.trace = json.load(f)
            except Exception as e:
                import logging
                logging.getLogger(__name__).error("Error loading trace file: %s", e)
        
    def generate_report(self):
        """
        Generates a Word document report of the transformation trace.
        Returns the BytesIO object of the document.
        """
        from io import BytesIO
        
        document = Document()
        
        # Style Configuration
        style = document.styles['Normal']
        font = style.font
        font.name = 'Calibri'
        font.size = Pt(11)
        
        # Title
        self._add_title(document)
        
        # Session Info
        self._add_session_info(document)
        
        # Steps
        document.add_heading('Transformation Steps', level=1)
        
        steps = self.trace.get("steps", [])
        if not steps:
            p = document.add_paragraph()
            p.add_run("No transformation steps recorded.").italic = True
        else:
            for i, step in enumerate(steps):
                self._add_step(document, i + 1, step)
                
        # Footer / Appendix?
        self._add_footer(document)
        
        # Save to buffer
        f = BytesIO()
        document.save(f)
        f.seek(0)
        return f

    def _add_title(self, doc):
        head = doc.add_heading('Data Transformation Report', 0)
        head.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(f"Generated on {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
        run.italic = True
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(128, 128, 128)

    def _add_session_info(self, doc):
        doc.add_heading('Session Information', level=1)
        
        table = doc.add_table(rows=0, cols=2)
        table.style = 'Table Grid'
        
        self._add_row(table, "Session ID", self.trace.get("session_id", "N/A"))
        self._add_row(table, "Source Dataset", self.trace.get("source_dataset", "N/A"))
        self._add_row(table, "Step Count", str(len(self.trace.get("steps", []))))

        doc.add_paragraph() # Spacer

    def _add_step(self, doc, index, step):
        # Step Header
        category = self._get_readable_function_name(step.get("function", "Unknown"))
        head = doc.add_heading(f"Step {index}: {category}", level=2)
        
        # Metadata line
        ts = step.get("timestamp", "")
        if ts:
            try:
                # Cleaning ISO format if needed
                ts = ts.replace("T", " ").split(".")[0]
            except:
                pass
        
        p_meta = doc.add_paragraph()
        run = p_meta.add_run(f"Executed at: {ts}")
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(100, 100, 100)
        
        # Description
        desc = step.get("description", "")
        if desc:
            doc.add_paragraph(desc)
            
        # Parameters Table
        params = step.get("params", {})
        if params:
            p = doc.add_paragraph()
            p.add_run("Parameters:").bold = True
            self._render_params(doc, params)
            
        doc.add_paragraph() # Spacer

    def _render_params(self, doc, params, indent=0):
        if not params:
            return

        table = doc.add_table(rows=0, cols=2)
        table.style = 'Light Shading' # Lighter style for params
        
        for k, v in params.items():
            key_cell = table.add_row().cells
            key_cell[0].text = str(k)
            
            # Format value
            if isinstance(v, list):
                if len(v) > 5 and isinstance(v[0], str):
                    key_cell[1].text = f"{len(v)} items: {', '.join(v[:5])}..."
                else:
                    key_cell[1].text = ", ".join([str(x) for x in v])
            elif isinstance(v, dict):
                 key_cell[1].text = json.dumps(v, indent=2)
            else:
                key_cell[1].text = str(v)

    def _add_row(self, table, label, value):
        row = table.add_row()
        row.cells[0].text = label
        row.cells[0].paragraphs[0].runs[0].font.bold = True
        row.cells[1].text = str(value)

    def _get_readable_function_name(self, func_name):
        mapping = {
            "initial_load": "Dataset Loading",
            "enrichment": "Data Enrichment",
            "imputation": "Missing Value Imputation",
            "variable_transformation": "Variable Transformation",
            "variable_computation": "New Variable Computation",
            "outlier_handling": "Outlier Detection & Handling",
            "data_type_conversion": "Data Type Conversion",
            "missing_value_handling": "Manual Missing Value Handling"
        }
        return mapping.get(func_name, func_name.replace("_", " ").title())

    def _add_footer(self, doc):
        pass
